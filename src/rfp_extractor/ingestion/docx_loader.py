"""DOCX ingestion (assignment briefs, specification documents)."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from ..models.document import Block, BlockKind, Document, FileFormat, Page
from .base import DocumentLoader, make_doc_id
from .normalize import normalize_text, squish

_ROW_SPLIT_RE = re.compile(r"\s*\|\s*")


class DocxLoader(DocumentLoader):
    formats = (FileFormat.DOCX,)

    def load(self, path: str | Path) -> Document:
        from docx import Document as DocxDocument  # imported lazily; optional dep

        path = Path(path)
        docx = DocxDocument(str(path))
        blocks: list[Block] = []
        label_values: dict[str, str] = {}
        n = 0

        for para in docx.paragraphs:
            text = squish(normalize_text(para.text))
            if not text:
                continue
            n += 1
            style = (para.style.name or "").lower() if para.style is not None else ""
            kind = BlockKind.HEADING if "heading" in style else BlockKind.TEXT
            block = Block(block_id=f"p1.b{n}", page_no=1, kind=kind, text=text, raw_text=para.text)
            m = re.match(r"^(?P<label>[A-Za-z][\w /'&.,#()-]{2,60}?)\s*:\s*(?P<value>.+)$", text)
            if m:
                block.label = m.group("label").strip()
                block.kind = BlockKind.LABEL_VALUE
                block.meta["label_value"] = m.group("value").strip()
                label_values.setdefault(block.label, block.meta["label_value"])
            blocks.append(block)

        n += 1
        for table in docx.tables:
            rows: list[list[str]] = []
            for row in table.rows:
                rows.append([squish(normalize_text(c.text)) for c in row.cells])
            if len(rows) < 2:
                continue
            for row in rows[1:]:
                # Keep rows whose *value* is empty too: a specification template
                # lists its field names with blank values, and those names are
                # exactly what we want to capture.
                if len(row) >= 2 and row[0]:
                    label_values.setdefault(row[0], row[1] if len(row) > 1 else "")
            text = "\n".join(" | ".join(r) for r in rows)
            blocks.append(
                Block(
                    block_id=f"p1.t{n}",
                    page_no=1,
                    kind=BlockKind.TABLE,
                    text=text,
                    raw_text=text,
                    table=rows,
                )
            )
            n += 1

        content_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        return Document(
            doc_id=make_doc_id(path.name, content_hash),
            file_name=path.name,
            source_path=str(path),
            file_format=FileFormat.DOCX,
            content_hash=content_hash,
            pages=[Page(page_no=1, blocks=blocks)],
            label_values=label_values,
            meta={"loader": "DocxLoader"},
        )
